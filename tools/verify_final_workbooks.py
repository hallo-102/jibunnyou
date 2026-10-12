"""Validate fresh copies of four production workbooks and independent rank oracle."""
from pathlib import Path
import sys
import importlib
import shutil
import hashlib
import json
from copy import copy
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
m = importlib.import_module('1_keibayosou_conditional_rerank')
OUT = ROOT / 'reports' / datetime.now().strftime('workbooks_%Y%m%d_%H%M%S')

def compare(before, after):
    # Include empty styled cells, formulas, cell data types and all original dimensions.
    a, b = load_workbook(before), load_workbook(after)
    issues = []
    for table in ('_fonts', '_fills', '_borders', '_alignments', '_number_formats', '_protections'):
        old, new = getattr(a, table), getattr(b, table)
        if list(old) != list(new)[:len(old)]:
            issues.append(f'workbook:{table}')
    for ws in a:
        if ws.title not in b:
            issues.append(f'{ws.title}:missing')
            continue
        dest = b[ws.title]
        for row in ws:
            for c in row:
                d = dest[c.coordinate]
                if (c.value, c.data_type) != (d.value, d.data_type):
                    issues.append(f'{ws.title}!{c.coordinate}:value/type')
                if c._style != d._style and any(copy(getattr(c, x)) != copy(getattr(d, x)) for x in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection')):
                    issues.append(f'{ws.title}!{c.coordinate}:style')
        for kind in ('column_dimensions', 'row_dimensions'):
            old, new = getattr(ws, kind), getattr(dest, kind)
            for key, dim in old.items():
                if key not in new or dict(dim) != dict(new[key]):
                    issues.append(f'{ws.title}:{kind}:{key}')
        if str(ws.merged_cells) != str(dest.merged_cells) or ws.freeze_panes != dest.freeze_panes:
            issues.append(f'{ws.title}:merge/freeze')
    a.close()
    b.close()
    return issues

def oracle(path):
    # Independent expected-order calculation; never use race results.
    t = pd.read_excel(path, sheet_name='TARGET')
    n = pd.read_excel(path, sheet_name='今走レース情報')
    key = 'rid_str' if 'rid_str' in t else 'レースID'
    nkey = 'rid_str' if 'rid_str' in n else 'レースID'
    score = 'score_raw' if 'score_raw' in t else 'score'
    form = 'recent_form_score' if 'recent_form_score' in t else '近走状態スコア'
    classes = [c for c in ('クラス', 'レース種別', 'レース名') if c in n]
    changed, excluded, errors = 0, 0, []
    for rid, group in t.groupby(key):
        ordered = group.sort_values('rank')
        leader = ordered.iloc[0]
        info = ' '.join(n.loc[n[nkey] == rid, classes].fillna('').astype(str).to_numpy().flatten())
        expected = ordered['馬番'].tolist()
        if '未勝利' in info or '新馬' in info:
            excluded += 1
        elif info and pd.notna(leader[score]) and pd.notna(leader[form]):
            options = []
            for _, horse in ordered.head(5).iloc[1:].iterrows():
                if pd.notna(horse[score]) and pd.notna(horse[form]):
                    adjusted = float(horse[score]) + .3 * (float(horse[form]) - 50)
                    if 0 <= float(leader[score]) - float(horse[score]) <= 4 and float(horse[form]) - float(leader[form]) >= 10 and adjusted > float(leader[score]) + .3 * (float(leader[form]) - 50):
                        options.append((-adjusted, horse['rank'], horse['馬番']))
            if options:
                chosen = sorted(options)[0][2]
                expected.remove(chosen)
                expected.insert(0, chosen)
                changed += 1
        actual = group.sort_values('条件付き順位')['馬番'].tolist()
        if actual != expected or sorted(group['条件付き順位']) != list(range(1, len(group) + 1)) or set(actual[:5]) != set(ordered['馬番'].tolist()[:5]):
            errors.append(str(rid))
    return {'races': t[key].nunique(), 'changed': changed, 'excluded': excluded, 'oracle_errors': errors}

def main():
    OUT.mkdir()
    results = []
    for date in ('20261003', '20261010', '20261011', '20261012'):
        source = ROOT.parent / 'keiba_yosou_2026-2/data/output' / f'馬の競走成績_with_feat_{date}.xlsx'
        dest = OUT / source.name
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        shutil.copy2(source, dest)
        count = m.append_conditional_ranking_sheet(str(dest))
        issues = compare(source, dest)
        stats = oracle(dest)
        once = OUT / f'once_{date}.xlsx'
        shutil.copy2(dest, once)
        m.append_conditional_ranking_sheet(str(dest))
        repeat = compare(once, dest)
        wb = load_workbook(dest, read_only=True)
        headers = [c.value for c in next(wb['TARGET'].rows)]
        duplicate = any(headers.count(c) != 1 for c in ('条件付き順位', '条件付き順位変更', '条件付き変更理由'))
        wb.close()
        result = {'date': date, **stats, 'reported_changes': count, 'preservation_issues': issues, 'repeat_issues': repeat, 'duplicate_columns': duplicate, 'source_unchanged': hashlib.sha256(source.read_bytes()).hexdigest() == original_hash}
        results.append(result)
        print(json.dumps({k: len(v) if isinstance(v, list) else v for k, v in result.items()}, ensure_ascii=False), flush=True)
    (ROOT / 'reports/workbook_results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')

if __name__ == '__main__':
    main()
