from __future__ import annotations

import argparse
import io
import json
import math
import re
import sqlite3
import statistics
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36"
DEFAULT_RACE_DATE = "20260922"
COURSE_CODES = {
    "01": "札幌",
    "02": "函館",
    "03": "福島",
    "04": "新潟",
    "05": "東京",
    "06": "中山",
    "07": "中京",
    "08": "京都",
    "09": "阪神",
    "10": "小倉",
}
KNOWN_MEETING_DAYS = {
    "20260919": "05",
    "20260920": "06",
    "20260921": "07",
    "20260922": "07",
}


def jra_schedule_url(race_date: str) -> str:
    """対象日のJRA公式競馬番組URLを組み立てる。"""
    parsed = datetime.strptime(race_date, "%Y%m%d")
    return f"https://www.jra.go.jp/keiba/calendar{parsed.year}/{parsed.year}/{parsed.month}/{parsed.strftime('%m%d')}.html"


def official_race_number_from_row(row: Any) -> int | None:
    """番組表の正式なR番号セルだけを解析し、中止注記内の番号を除外する。"""
    first_cell = row.find(["th", "td"])
    if first_cell is None:
        return None
    match = re.fullmatch(r"\s*(\d+)\s*レース\s*", first_cell.get_text(" ", strip=True))
    return int(match.group(1)) if match else None


def official_race_counts_from_html(html: bytes | str) -> dict[str, int]:
    """JRA公式番組表から、開催場別に実施予定のレース数を抽出する。"""
    soup = BeautifulSoup(html, "html.parser")
    counts: dict[str, int] = {}
    for table in soup.select("table.basic"):
        heading = table.select_one("div.main")
        heading_text = heading.get_text(" ", strip=True) if heading else table.get_text(" ", strip=True)[:40]
        racecourse = next((name for name in COURSE_CODES.values() if name in heading_text), None)
        if racecourse is None:
            continue
        race_numbers: set[int] = set()
        for row in table.select("tr"):
            race_number = official_race_number_from_row(row)
            if race_number is not None:
                race_numbers.add(race_number)
        counts[racecourse] = max(counts.get(racecourse, 0), len(race_numbers))
    return counts


def official_race_ids_from_html(html: bytes | str, race_date: str) -> list[str]:
    """JRA公式番組表の開催回・開催日・R番号から12桁レースIDを復元する。"""
    soup = BeautifulSoup(html, "html.parser")
    course_to_code = {name: code for code, name in COURSE_CODES.items()}
    race_ids: list[str] = []
    for table in soup.select("table.basic"):
        heading = table.select_one("div.main")
        heading_text = heading.get_text(" ", strip=True) if heading else ""
        racecourse = next((name for name in COURSE_CODES.values() if name in heading_text), None)
        if racecourse is None:
            continue
        meeting_match = re.search(r"(\d+)\s*回", heading_text)
        day_match = re.search(r"(\d+)\s*日", heading_text)
        if meeting_match is None or day_match is None:
            continue
        meeting_no = int(meeting_match.group(1))
        meeting_day = int(day_match.group(1))
        for row in table.select("tr"):
            race_number = official_race_number_from_row(row)
            if race_number is not None:
                race_ids.append(
                    f"{race_date[:4]}{course_to_code[racecourse]}{meeting_no:02d}{meeting_day:02d}{race_number:02d}"
                )
    return sorted(set(race_ids))


def fetch_official_race_program(race_date: str) -> tuple[dict[str, int], list[str]]:
    """JRA公式番組表を一度取得し、開催場別件数と正規レースIDを返す。"""
    response = requests.get(jra_schedule_url(race_date), headers={"User-Agent": USER_AGENT}, timeout=30)
    response.raise_for_status()
    counts = official_race_counts_from_html(response.content)
    race_ids = official_race_ids_from_html(response.content, race_date)
    if not counts or not race_ids:
        raise RuntimeError(f"JRA公式競馬番組を解析できません: {race_date}")
    verify_official_race_counts(race_ids, counts)
    return counts, race_ids


def fetch_official_race_counts(race_date: str) -> dict[str, int]:
    """JRA公式番組表を取得し、解析不能時は安全側で処理を停止する。"""
    counts, _ = fetch_official_race_program(race_date)
    return counts


def verify_official_race_counts(race_ids: list[str], official_counts: dict[str, int]) -> None:
    """保存前に、取得対象のレース数がJRA公式番組表と一致することを保証する。"""
    discovered: Counter[str] = Counter()
    for race_id in race_ids:
        code = str(race_id)[4:6]
        if code not in COURSE_CODES:
            raise RuntimeError(f"未対応のJRA競馬場コードです: {code} ({race_id})")
        discovered[COURSE_CODES[code]] += 1
    expected = {course: count for course, count in official_counts.items() if count > 0}
    actual = {course: count for course, count in discovered.items() if count > 0}
    if actual != expected:
        raise RuntimeError(
            "JRA公式レース数と不一致のため保存を中止します: "
            f"official={json.dumps(expected, ensure_ascii=False, sort_keys=True)} "
            f"discovered={json.dumps(actual, ensure_ascii=False, sort_keys=True)}"
        )


def racecourse_from_source_id(source_race_id: str) -> str:
    """JRAレースIDの競馬場コードを競馬場名へ変換する。"""
    return COURSE_CODES.get(str(source_race_id)[4:6], "データなし")


def clean_column(value: object) -> str:
    """表の列名に含まれる空白を除去し、安定した参照名へ正規化する。"""
    return re.sub(r"\s+", "", str(value))


def parse_time_seconds(value: object) -> float | None:
    """競走時計の文字列を秒へ変換し、中止などの非数値はNULL相当にする。"""
    text = str(value).strip()
    match = re.fullmatch(r"(?:(\d+):)?(\d{1,2})\.(\d)", text)
    if not match:
        return None
    minutes = int(match.group(1) or 0)
    return minutes * 60 + int(match.group(2)) + int(match.group(3)) / 10


def parse_class(text: str) -> str:
    """レース情報から比較用のクラス区分を取り出す。"""
    for label, pattern in (
        ("新馬", r"新馬"),
        ("未勝利", r"未勝利"),
        ("1勝", r"１勝|1勝"),
        ("2勝", r"２勝|2勝"),
        ("3勝", r"３勝|3勝"),
        ("オープン", r"オープン|G[ⅠⅡⅢ123]|重賞"),
    ):
        if re.search(pattern, text):
            return label
    return "その他"


def parse_going(text: str) -> str | None:
    """芝・ダートの馬場状態を抽出する。障害で二つある場合は先頭を採用する。"""
    match = re.search(r"馬場\s*:\s*(稍重|不良|良|重|稍|不)", text)
    if not match:
        return None
    return {"不": "不良", "稍": "稍重"}.get(match.group(1), match.group(1))


def parse_surface_distance(text: str) -> tuple[str | None, int | None]:
    """レース情報から馬場種別と距離を抽出する。"""
    match = re.search(r"/\s*(芝|ダート|ダ|障)\s*(\d{3,4})m", text)
    if not match:
        return None, None
    surface = {"芝": "芝", "ダ": "ダート", "ダート": "ダート", "障": "障害"}[match.group(1)]
    return surface, int(match.group(2))


def distance_group(distance: int | None) -> str:
    """距離を短距離・マイル・中距離・長距離へ分類する。"""
    if distance is None:
        return "データなし"
    if distance <= 1400:
        return "短距離"
    if distance <= 1700:
        return "マイル"
    if distance <= 2200:
        return "中距離"
    return "長距離"


def load_historical_clock_reference(path: Path) -> pd.DataFrame:
    """既存結果ブックを読み、同条件時計比較に必要な勝ち馬行だけを抽出する。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
        except StopIteration:
            continue
        index = {name: pos for pos, name in enumerate(headers)}
        required = {"レースID", "着 順", "タイム", "レース情報"}
        if not required.issubset(index):
            continue
        for values in iterator:
            finish = values[index["着 順"]]
            if str(finish).strip() not in {"1", "1.0"}:
                continue
            race_id = str(values[index["レースID"]] or "").strip()
            race_info = str(values[index["レース情報"]] or "")
            if len(race_id) < 6 or race_id[4:6] not in COURSE_CODES:
                continue
            surface, distance = parse_surface_distance(race_info)
            going = parse_going(race_info)
            seconds = parse_time_seconds(values[index["タイム"]])
            if surface not in {"芝", "ダート"} or distance is None or seconds is None:
                continue
            rows.append(
                {
                    "race_id": race_id,
                    "racecourse": racecourse_from_source_id(race_id),
                    "surface": surface,
                    "distance": distance,
                    "going": going,
                    "class_name": parse_class(race_info),
                    "winner_time_sec": seconds,
                }
            )
    workbook.close()
    return pd.DataFrame(rows)


def discover_race_ids(project_root: Path, history_path: Path, race_date: str) -> list[str]:
    """当日出馬表を優先し、なければ保存済み結果シートから実施レースIDを復元する。"""
    entries_path = project_root / "keiba_platform_v2" / "data" / "raw" / f"netkeiba_entries_{race_date}.csv"
    if entries_path.exists():
        entries = pd.read_csv(entries_path, encoding="utf-8-sig")
        if "source_race_id" in entries.columns:
            race_ids = entries["source_race_id"].dropna().astype(str).str.replace(r"\.0$", "", regex=True)
            return sorted(race_ids.drop_duplicates().tolist())

    workbook = load_workbook(history_path, read_only=True, data_only=True)
    if race_date not in workbook.sheetnames:
        workbook.close()
        raise RuntimeError(f"結果シートがありません: {race_date}")
    sheet = workbook[race_date]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
    index = {name: pos for pos, name in enumerate(headers)}
    if not {"レースID", "着 順"}.issubset(index):
        workbook.close()
        raise RuntimeError(f"結果シートの必須列がありません: {race_date}")
    race_ids: set[str] = set()
    for values in iterator:
        finish = pd.to_numeric(values[index["着 順"]], errors="coerce")
        raw_race_id = values[index["レースID"]]
        if pd.isna(finish) or raw_race_id is None:
            continue
        race_id = re.sub(r"\.0$", "", str(raw_race_id).strip())
        expected_day = KNOWN_MEETING_DAYS.get(race_date)
        if (
            len(race_id) == 12
            and race_id[4:6] in COURSE_CODES
            and (expected_day is None or race_id[8:10] == expected_day)
        ):
            race_ids.add(race_id)
    workbook.close()
    return sorted(race_ids)


def load_entry_horse_id_map(project_root: Path, race_date: str) -> dict[tuple[str, int], str]:
    """事前出馬表に保存されたhorse_idをレースID・馬番単位で読み込む。"""
    candidates = (
        project_root / "keiba_platform_v2" / "data" / "raw" / f"netkeiba_entries_{race_date}.csv",
        project_root / "keiba_platform_v2" / "data" / "input" / f"races_{race_date}.csv",
    )
    mapping: dict[tuple[str, int], str] = {}
    for path in candidates:
        if not path.exists():
            continue
        frame = pd.read_csv(path, encoding="utf-8-sig", dtype={"source_race_id": str, "horse_id": str})
        required = {"source_race_id", "horse_no", "horse_id"}
        if not required.issubset(frame.columns):
            continue
        for row in frame.itertuples(index=False):
            race_id = re.sub(r"\.0$", "", str(getattr(row, "source_race_id", "")).strip())
            horse_id = re.sub(r"\.0$", "", str(getattr(row, "horse_id", "")).strip())
            horse_no = pd.to_numeric(getattr(row, "horse_no", None), errors="coerce")
            if race_id and horse_id and horse_id.lower() != "nan" and pd.notna(horse_no):
                mapping[(race_id, int(horse_no))] = horse_id
        if mapping:
            break
    return mapping


def result_horse_ids_by_name(soup: BeautifulSoup) -> dict[str, str]:
    """結果表の競走馬リンクからhorse_idを抽出し、馬名へ対応付ける。"""
    mapping: dict[str, str] = {}
    for anchor in soup.select('table.RaceTable01 a[href*="/horse/"]'):
        match = re.search(r"/horse/(\d+)", str(anchor.get("href", "")))
        horse_name = anchor.get_text(" ", strip=True)
        if match and horse_name:
            mapping[horse_name] = match.group(1)
    return mapping


def historical_standard(reference: pd.DataFrame, race: dict[str, Any]) -> tuple[float | None, int, str]:
    """同競馬場・距離・馬場・クラスを優先し、十分な標本がある中央値を返す。"""
    base = reference[
        reference["racecourse"].eq(race["racecourse"])
        & reference["surface"].eq(race["surface"])
        & reference["distance"].eq(race["distance"])
        & reference["race_id"].astype(str).str[:10].lt(race["source_race_id"][:10])
    ]
    candidates = (
        (base[base["going"].eq(race["going"]) & base["class_name"].eq(race["class_name"])], "距離×馬場×クラス"),
        (base[base["going"].eq(race["going"])], "距離×馬場"),
        (base[base["class_name"].eq(race["class_name"])], "距離×クラス"),
        (base, "距離"),
    )
    for frame, basis in candidates:
        values = frame["winner_time_sec"].dropna().astype(float).tolist()
        if len(values) >= 3:
            return float(statistics.median(values)), len(values), basis
    return None, 0, "比較標本不足"


def normalize_result_table(table: pd.DataFrame) -> pd.DataFrame:
    """netkeiba結果表の現在の列名を回顧用の列名へ変換する。"""
    frame = table.copy()
    frame.columns = [clean_column(c) for c in frame.columns]
    aliases = {
        "着順": "finish_raw",
        "枠": "frame_no",
        "馬番": "horse_no",
        "馬名": "horse_name",
        "性齢": "sex_age",
        "斤量": "carried_weight",
        "騎手": "jockey",
        "タイム": "finish_time",
        "着差": "margin",
        "人気": "popularity",
        "単勝オッズ": "odds",
        "後3F": "last3f",
        "コーナー通過順": "passing_order",
        "厩舎": "trainer",
        "馬体重(増減)": "body_weight",
    }
    frame = frame.rename(columns={c: aliases.get(c, c) for c in frame.columns})
    return frame


def parse_laps(tables: list[pd.DataFrame]) -> tuple[list[float], list[str]]:
    """ラップ表から各200mラップと累積時計を取り出す。"""
    for table in tables:
        if isinstance(table.columns, pd.Index) and "200m" in [str(c) for c in table.columns] and len(table) >= 2:
            cumulative = [str(v) for v in table.iloc[0].tolist()]
            lap_values = [pd.to_numeric(v, errors="coerce") for v in table.iloc[1].tolist()]
            laps = [float(v) for v in lap_values if pd.notna(v)]
            return laps, cumulative
    return [], []


def classify_pace(label: str | None, laps: list[float]) -> tuple[str, float | None, float | None]:
    """提供元のH/M/Sと前後3F差を組み合わせて5段階のペースを記録する。"""
    if len(laps) < 6:
        return "データなし", None, None
    early = round(sum(laps[:3]), 1)
    late = round(sum(laps[-3:]), 1)
    difference = early - late
    if label == "H":
        pace = "超ハイ" if difference <= -2.0 else "ハイ"
    elif label == "S":
        pace = "超スロー" if difference >= 2.0 else "スロー"
    else:
        pace = "平均"
    return pace, early, late


def last_corner_position(value: object) -> int | None:
    """各馬の通過順位から最終コーナー位置を取得する。"""
    numbers = re.findall(r"\d+", str(value))
    return int(numbers[-1]) if numbers else None


def running_style(position: int | None, field_size: int) -> str:
    """最終コーナー位置を頭数比で脚質区分へ変換する。"""
    if position is None or field_size <= 0:
        return "データなし"
    ratio = position / field_size
    if position == 1:
        return "逃げ"
    if ratio <= 0.25:
        return "先行"
    if ratio <= 0.60:
        return "中団"
    if ratio <= 0.85:
        return "差し"
    return "追込"


def payout_summary(tables: list[pd.DataFrame]) -> str:
    """払戻表を改変せず、検索可能なJSON文字列へまとめる。"""
    payouts: dict[str, dict[str, str]] = {}
    for table in tables[1:3]:
        for _, row in table.iterrows():
            values = [str(v).replace("nan", "").strip() for v in row.tolist()]
            if not values or not values[0]:
                continue
            payouts[values[0]] = {
                "selection": values[1] if len(values) > 1 else "",
                "payout": values[2] if len(values) > 2 else "",
                "popularity": values[3] if len(values) > 3 else "",
            }
    return json.dumps(payouts, ensure_ascii=False)


def scrape_race(
    source_race_id: str,
    race_date: str,
    reference: pd.DataFrame,
    entry_horse_ids: dict[tuple[str, int], str] | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """1レース分の結果・メタデータ・各馬データをnetkeibaから取得する。"""
    url = f"https://race.netkeiba.com/race/result.html?race_id={source_race_id}&rf=race_list"
    response = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=30)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    tables = pd.read_html(io.BytesIO(response.content))
    if not tables:
        raise RuntimeError(f"結果表がありません: {source_race_id}")

    race_data_01 = soup.select_one(".RaceData01")
    race_data_02 = soup.select_one(".RaceData02")
    race_name_node = soup.select_one(".RaceName")
    meta_01 = race_data_01.get_text(" ", strip=True) if race_data_01 else ""
    meta_02 = race_data_02.get_text(" ", strip=True) if race_data_02 else ""
    race_name = race_name_node.get_text(" ", strip=True) if race_name_node else "データなし"
    surface, distance = parse_surface_distance("/" + meta_01.lstrip("/"))
    going = parse_going(meta_01)
    start_match = re.search(r"(\d{1,2}:\d{2})発走", meta_01)
    weather_match = re.search(r"天候\s*:\s*([^/\s]+)", meta_01)
    field_match = re.search(r"(\d+)頭", meta_02)
    class_name = parse_class(meta_02 + " " + race_name)
    race_no = int(source_race_id[-2:])
    racecourse = racecourse_from_source_id(source_race_id)
    field_size = int(field_match.group(1)) if field_match else len(tables[0])
    result = normalize_result_table(tables[0])
    result_horse_ids = result_horse_ids_by_name(soup)
    laps, cumulative = parse_laps(tables)
    pace_node = soup.select_one(".RapPace_Title")
    pace_match = re.search(r"ペース\s*:\s*([HMS])", pace_node.get_text(" ", strip=True) if pace_node else "")
    pace_type, early3f, late3f = classify_pace(pace_match.group(1) if pace_match else None, laps)

    winner_time = parse_time_seconds(result.iloc[0].get("finish_time")) if not result.empty else None
    race_stub = {"source_race_id": source_race_id, "racecourse": racecourse, "surface": surface, "distance": distance, "going": going, "class_name": class_name}
    standard, standard_n, standard_basis = historical_standard(reference, race_stub) if surface in {"芝", "ダート"} else (None, 0, "障害除外")
    clock_diff = round(winner_time - standard, 2) if winner_time is not None and standard is not None else None
    clock_diff_per_1000m = (
        round(clock_diff * 1000 / distance, 4)
        if clock_diff is not None and distance
        else None
    )

    horse_rows: list[dict[str, Any]] = []
    for _, row in result.iterrows():
        finish_text = str(row.get("finish_raw", ""))
        finish_match = re.search(r"\d+", finish_text)
        finish = int(finish_match.group()) if finish_match else None
        horse_no_value = pd.to_numeric(row.get("horse_no"), errors="coerce")
        popularity_value = pd.to_numeric(row.get("popularity"), errors="coerce")
        odds_value = pd.to_numeric(row.get("odds"), errors="coerce")
        last3f_value = pd.to_numeric(row.get("last3f"), errors="coerce")
        position = last_corner_position(row.get("passing_order"))
        horse_name = str(row.get("horse_name", "")).strip()
        horse_no = int(horse_no_value) if pd.notna(horse_no_value) else None
        horse_id = result_horse_ids.get(horse_name)
        if horse_id is None and horse_no is not None and entry_horse_ids is not None:
            horse_id = entry_horse_ids.get((source_race_id, horse_no))
        horse_rows.append(
            {
                "date": race_date,
                "racecourse": racecourse,
                "race_no": race_no,
                "source_race_id": source_race_id,
                "horse_id": horse_id,
                "horse_name": horse_name,
                "horse_no": horse_no,
                "frame_no": int(pd.to_numeric(row.get("frame_no"), errors="coerce")) if pd.notna(pd.to_numeric(row.get("frame_no"), errors="coerce")) else None,
                "finish": finish,
                "finish_status": finish_text,
                "popularity": int(popularity_value) if pd.notna(popularity_value) else None,
                "odds": float(odds_value) if pd.notna(odds_value) else None,
                "finish_time": str(row.get("finish_time", "")),
                "margin": str(row.get("margin", "")),
                "last3f": float(last3f_value) if pd.notna(last3f_value) else None,
                "passing_order": str(row.get("passing_order", "")),
                "last_corner_position": position,
                "running_style": running_style(position, field_size),
                "carried_weight": float(pd.to_numeric(row.get("carried_weight"), errors="coerce")) if pd.notna(pd.to_numeric(row.get("carried_weight"), errors="coerce")) else None,
                "jockey": str(row.get("jockey", "")),
                "trainer": str(row.get("trainer", "")),
                "body_weight": str(row.get("body_weight", "")),
                "prediction_rank": None,
                "prediction_score": None,
                "actual_performance": "B",
                "trip_trouble": "データなし",
                "pace_advantage": 0,
                "pace_disadvantage": 0,
                "track_bias_advantage": None,
                "track_bias_disadvantage": None,
                "inside_outside_advantage": None,
                "distance_loss": None,
                "late_speed_rating": None,
                "opponent_strength": None,
                "race_level": "データなし",
                "next_start_watch_score": 0,
                "next_start_watch_reason": "",
                "next_start_target_condition": "",
                "next_start_risk": "",
            }
        )

    valid_late = sorted([h for h in horse_rows if h["last3f"] is not None], key=lambda h: h["last3f"])
    for rank, horse in enumerate(valid_late, start=1):
        horse["late_speed_rating"] = rank
    for horse in horse_rows:
        style = horse["running_style"]
        front = style in {"逃げ", "先行"}
        closer = style in {"差し", "追込"}
        slow = pace_type in {"スロー", "超スロー"}
        high = pace_type in {"ハイ", "超ハイ"}
        horse["pace_advantage"] = int((slow and front) or (high and closer))
        horse["pace_disadvantage"] = int((high and front) or (slow and closer))

        reasons: list[str] = []
        score = 0
        if horse["finish"] is not None and 4 <= horse["finish"] <= 12 and horse["late_speed_rating"] is not None and horse["late_speed_rating"] <= 3:
            score += 40
            reasons.append(f"上がり{horse['late_speed_rating']}位")
        if horse["finish"] is not None and horse["finish"] >= 4 and horse["pace_disadvantage"]:
            score += 25
            reasons.append(f"{pace_type}で{style}に展開不利")
        if horse["popularity"] is not None and horse["popularity"] >= 6 and score > 0:
            score += 15
            reasons.append("市場評価が低い")
        horse["next_start_watch_score"] = min(score, 100)
        horse["next_start_watch_reason"] = "、".join(reasons)
        if score >= 50:
            horse["actual_performance"] = "A"
            horse["next_start_target_condition"] = "今回より展開適合度が高い条件"
        elif horse["finish"] is not None and horse["finish"] <= 3 and horse["pace_advantage"]:
            horse["actual_performance"] = "D" if pace_type in {"超スロー", "超ハイ"} else "C"
            horse["next_start_risk"] = "今回と逆のペースになった場合"

    ordered = sorted([h for h in horse_rows if h["finish"] is not None], key=lambda h: h["finish"])
    race = {
        "date": race_date,
        "racecourse": racecourse,
        "race_no": race_no,
        "source_race_id": source_race_id,
        "race_name": race_name,
        "surface": surface,
        "distance": distance,
        "distance_group": distance_group(distance),
        "going": going,
        "class_name": class_name,
        "field_size": field_size,
        "weather": weather_match.group(1) if weather_match else None,
        "start_time": start_match.group(1) if start_match else None,
        "winner": ordered[0]["horse_name"] if len(ordered) > 0 else None,
        "second": ordered[1]["horse_name"] if len(ordered) > 1 else None,
        "third": ordered[2]["horse_name"] if len(ordered) > 2 else None,
        "winning_time": ordered[0]["finish_time"] if len(ordered) > 0 else None,
        "pace_type": pace_type,
        "early3f": early3f,
        "late3f_race": late3f,
        "first1000m": cumulative[4] if len(cumulative) >= 5 else None,
        "lap_times": json.dumps(laps, ensure_ascii=False),
        "race_level": "データなし",
        "clock_rating": None,
        "clock_diff_sec": clock_diff,
        "clock_diff_per_1000m": clock_diff_per_1000m,
        "clock_reference_sec": standard,
        "clock_reference_n": standard_n,
        "clock_reference_basis": standard_basis,
        "pace_rating": pace_type,
        "inside_outside_bias": "データなし",
        "running_style_bias": None,
        "prediction_top1": None,
        "prediction_top1_finish": None,
        "prediction_top3_hit": None,
        "prediction_top5_hit": None,
        "winner_prediction_rank": None,
        "payout": payout_summary(tables),
        "review_summary": "",
        "prediction_error_type": "当日予想未生成（運用処理未完了）",
        "bet_or_skip_review": "評価不能（事前予想なし）",
        "source_url": url,
    }
    time.sleep(0.2)
    return race, horse_rows


def clock_label(diff_sec: float | None, distance: int | None) -> str:
    """距離当たりの標準差を5段階の時計評価へ変換する。"""
    if diff_sec is None or not distance:
        return "データなし"
    return clock_label_normalized(diff_sec * 1000 / distance)


def clock_label_normalized(normalized: float | None) -> str:
    """1000m当たりの時計差から、距離に依存しない5段階評価を返す。"""
    if normalized is None:
        return "データなし"
    if normalized <= -1.0:
        return "非常に高速"
    if normalized <= -0.35:
        return "高速"
    if normalized < 0.35:
        return "標準"
    if normalized < 1.0:
        return "やや時計が掛かる"
    return "時計が掛かる"


def build_track_bias(races: list[dict[str, Any]], horses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """芝とダートを分離し、当日の時計・脚質傾向を集計する。"""
    rows: list[dict[str, Any]] = []
    race_date = races[0]["date"] if races else ""
    for racecourse in sorted({r["racecourse"] for r in races}):
      for surface in ("芝", "ダート"):
        surface_races = [r for r in races if r["racecourse"] == racecourse and r["surface"] == surface]
        if not surface_races:
            continue
        race_keys = {(r["racecourse"], r["race_no"]) for r in surface_races}
        surface_horses = [h for h in horses if (h["racecourse"], h["race_no"]) in race_keys]
        diffs = [float(r["clock_diff_sec"]) for r in surface_races if r["clock_diff_sec"] is not None]
        mean_diff = round(statistics.mean(diffs), 2) if diffs else None
        normalized_diffs = [
            float(r["clock_diff_per_1000m"])
            for r in surface_races
            if r.get("clock_diff_per_1000m") is not None
        ]
        mean_normalized = round(statistics.mean(normalized_diffs), 4) if normalized_diffs else None

        style_stats: dict[str, dict[str, float]] = {}
        for style in ("逃げ", "先行", "中団", "差し", "追込"):
            members = [h for h in surface_horses if h["running_style"] == style and h["finish"] is not None]
            wins = sum(1 for h in members if h["finish"] == 1)
            places = sum(1 for h in members if h["finish"] <= 3)
            style_stats[style] = {
                "starts": len(members),
                "wins": wins,
                "places": places,
                "win_rate": wins / len(members) if members else 0.0,
                "place_rate": places / len(members) if members else 0.0,
            }
        front_starts = style_stats["逃げ"]["starts"] + style_stats["先行"]["starts"]
        front_wins = style_stats["逃げ"]["wins"] + style_stats["先行"]["wins"]
        front_places = style_stats["逃げ"]["places"] + style_stats["先行"]["places"]
        closer_starts = style_stats["差し"]["starts"] + style_stats["追込"]["starts"]
        closer_wins = style_stats["差し"]["wins"] + style_stats["追込"]["wins"]
        closer_places = style_stats["差し"]["places"] + style_stats["追込"]["places"]
        front_place_rate = front_places / front_starts if front_starts else 0.0
        closer_place_rate = closer_places / closer_starts if closer_starts else 0.0
        style_gap = front_place_rate - closer_place_rate
        if style_gap >= 0.12:
            style_bias, style_score = "強い逃げ先行有利", 2
        elif style_gap >= 0.05:
            style_bias, style_score = "やや先行有利", 1
        elif style_gap <= -0.12:
            style_bias, style_score = "強い差し追込有利", -2
        elif style_gap <= -0.05:
            style_bias, style_score = "やや差し有利", -1
        else:
            style_bias, style_score = "フラット", 0

        segment_summaries: list[str] = []
        for label, low, high in (("前半", 1, 4), ("中盤", 5, 8), ("後半", 9, 12)):
            winners = [h for h in surface_horses if low <= h["race_no"] <= high and h["finish"] == 1]
            front_count = sum(1 for h in winners if h["running_style"] in {"逃げ", "先行"})
            closer_count = sum(1 for h in winners if h["running_style"] in {"差し", "追込"})
            summary = f"{label}:前{front_count}/差追{closer_count}/対象{len(winners)}"
            segment_summaries.append(summary)

        confidence = "MEDIUM" if len(surface_races) >= 5 else "LOW"
        rows.append(
            {
                "date": race_date,
                "racecourse": racecourse,
                "surface": surface,
                "distance_group": "全距離",
                "going": "/".join(sorted({str(r["going"]) for r in surface_races if r["going"]})),
                "race_count": len(surface_races),
                "clock_bias": clock_label_normalized(mean_normalized),
                "clock_bias_score": None if mean_normalized is None else round(-mean_normalized, 4),
                "clock_diff_sec": mean_diff,
                "clock_diff_per_1000m": mean_normalized,
                "inside_outside_bias": "データなし",
                "inside_outside_score": None,
                "pace_style_bias": style_bias,
                "pace_style_score": style_score,
                "front_runner_score": round(style_stats["逃げ"]["place_rate"] * 100, 1),
                "stalker_score": round(style_stats["先行"]["place_rate"] * 100, 1),
                "midpack_score": round(style_stats["中団"]["place_rate"] * 100, 1),
                "closer_score": round(style_stats["差し"]["place_rate"] * 100, 1),
                "deep_closer_score": round(style_stats["追込"]["place_rate"] * 100, 1),
                "inside_win_rate": None,
                "outside_win_rate": None,
                "front_win_rate": round(front_wins / front_starts, 4) if front_starts else None,
                "front_place_rate": round(front_place_rate, 4) if front_starts else None,
                "closer_win_rate": round(closer_wins / closer_starts, 4) if closer_starts else None,
                "closer_place_rate": round(closer_place_rate, 4) if closer_starts else None,
                "early_race_bias": segment_summaries[0],
                "middle_race_bias": segment_summaries[1],
                "late_race_bias": segment_summaries[2],
                "weather": "/".join(sorted({str(r["weather"]) for r in surface_races if r["weather"]})),
                "rainfall": None,
                "wind_direction": None,
                "wind_speed": None,
                "confidence": confidence,
                "notes": "実走進路データがないため内外判定は行わない。脚質は最終コーナー位置比から機械分類。",
                "style_stats_json": json.dumps(style_stats, ensure_ascii=False),
            }
        )
    return rows


def build_track_condition_stats(races: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """距離区分・馬場・クラス別に、1000m当たり時計差を集計する。"""
    grouped: dict[tuple[str, str, str, str, str, str], list[float]] = defaultdict(list)
    for race in races:
        if race.get("surface") not in {"芝", "ダート"}:
            continue
        value = race.get("clock_diff_per_1000m")
        if value is None:
            continue
        key = (
            str(race["date"]),
            str(race["racecourse"]),
            str(race["surface"]),
            str(race["distance_group"]),
            str(race.get("going") or "データなし"),
            str(race.get("class_name") or "データなし"),
        )
        grouped[key].append(float(value))

    rows: list[dict[str, Any]] = []
    for key, values in sorted(grouped.items()):
        race_date, racecourse, surface, group, going, class_name = key
        mean_value = round(statistics.mean(values), 4)
        median_value = round(statistics.median(values), 4)
        rows.append(
            {
                "date": race_date,
                "racecourse": racecourse,
                "surface": surface,
                "distance_group": group,
                "going": going,
                "class_name": class_name,
                "race_count": len(values),
                "clock_diff_per_1000m_mean": mean_value,
                "clock_diff_per_1000m_median": median_value,
                "clock_diff_per_1000m_std": round(statistics.stdev(values), 4) if len(values) >= 2 else None,
                "clock_diff_per_1000m_min": round(min(values), 4),
                "clock_diff_per_1000m_max": round(max(values), 4),
                "clock_bias": clock_label_normalized(mean_value),
                "confidence": "MEDIUM" if len(values) >= 5 else "LOW",
            }
        )
    return rows


def enrich_race_summaries(races: list[dict[str, Any]], horses: list[dict[str, Any]]) -> None:
    """全レースに、取得事実と不明点を分離した短い回顧文を付与する。"""
    for race in races:
        race["clock_rating"] = clock_label(race["clock_diff_sec"], race["distance"])
        participants = [
            h for h in horses
            if h["racecourse"] == race["racecourse"] and h["race_no"] == race["race_no"]
        ]
        top3_styles = [h["running_style"] for h in participants if h["finish"] is not None and h["finish"] <= 3]
        race["running_style_bias"] = "・".join(top3_styles) if top3_styles else "データなし"
        clock_text = f"標準比{race['clock_diff_sec']:+.2f}秒（{race['clock_reference_basis']} n={race['clock_reference_n']}）" if race["clock_diff_sec"] is not None else "時計比較は標本不足"
        race["review_summary"] = (
            f"{race['pace_type']}。1～3着脚質は{race['running_style_bias']}。{clock_text}。"
            "実走進路・接触・前詰まりは映像データ未取得のため評価しない。"
        )


def build_watch_rows(horses: list[dict[str, Any]], race_date: str) -> list[dict[str, Any]]:
    """事前に定義した定量ルールで次走注目候補を最大10頭抽出する。"""
    candidates = sorted(
        [h for h in horses if h["next_start_watch_score"] >= 50],
        key=lambda h: (-h["next_start_watch_score"], h["finish"] or 99),
    )[:10]
    valid_until = (datetime.strptime(race_date, "%Y%m%d") + timedelta(days=180)).strftime("%Y-%m-%d")
    rows: list[dict[str, Any]] = []
    for index, horse in enumerate(candidates, start=1):
        rows.append(
            {
                "date": race_date,
                "source_race_id": horse["source_race_id"],
                "horse_id": horse["horse_id"],
                "horse_name": horse["horse_name"],
                "racecourse": horse["racecourse"],
                "race_no": horse["race_no"],
                "finish": horse["finish"],
                "popularity": horse["popularity"],
                "watch_rank": index,
                "watch_reason": horse["next_start_watch_reason"],
                "trouble_type": "データなし",
                "hidden_strength": horse["next_start_watch_reason"],
                "preferred_surface": None,
                "preferred_distance": None,
                "preferred_course": None,
                "preferred_going": None,
                "expected_market_mispricing": "人気6番手以下" if (horse["popularity"] or 0) >= 6 else "不明",
                "avoid_if_favorite": 0,
                "valid_until": valid_until,
            }
        )
    return rows


def build_danger_rows(horses: list[dict[str, Any]], races: list[dict[str, Any]], race_date: str) -> list[dict[str, Any]]:
    """展開利が明示できる好走馬のみを次走危険候補として抽出する。"""
    race_lookup = {(r["racecourse"], r["race_no"]): r for r in races}
    valid_until = (datetime.strptime(race_date, "%Y%m%d") + timedelta(days=180)).strftime("%Y-%m-%d")
    rows: list[dict[str, Any]] = []
    for horse in horses:
        if horse["finish"] is None or horse["finish"] > 3 or not horse["pace_advantage"]:
            continue
        pace = race_lookup[(horse["racecourse"], horse["race_no"])]["pace_type"]
        rows.append(
            {
                "date": race_date,
                "source_race_id": horse["source_race_id"],
                "horse_id": horse["horse_id"],
                "horse_name": horse["horse_name"],
                "racecourse": horse["racecourse"],
                "race_no": horse["race_no"],
                "finish": horse["finish"],
                "popularity": horse["popularity"],
                "advantage_type": "ペース",
                "pace_advantage": 1,
                "track_bias_advantage": None,
                "trip_advantage": None,
                "danger_reason": f"{pace}と{horse['running_style']}の組合せに展開利",
                "danger_condition": "今回と逆方向のペースで人気する場合",
                "valid_until": valid_until,
            }
        )
    return rows


def build_horse_master_rows(horses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """horse_id単位で馬名と観測期間をまとめ、縦断結合用の馬マスタを作る。"""
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for horse in horses:
        if horse.get("horse_id"):
            grouped[str(horse["horse_id"])].append(horse)
    rows: list[dict[str, Any]] = []
    for horse_id, members in sorted(grouped.items()):
        ordered = sorted(members, key=lambda row: str(row["date"]))
        rows.append(
            {
                "horse_id": horse_id,
                "horse_name": ordered[-1]["horse_name"],
                "first_seen_date": ordered[0]["date"],
                "last_seen_date": ordered[-1]["date"],
            }
        )
    return rows


def calculate_next_start_outcomes(
    signals: list[dict[str, Any]],
    horse_history: list[dict[str, Any]],
    as_of_date: str,
) -> list[dict[str, Any]]:
    """horse_idで各シグナル後の最初の完走レースを結び、単勝100円換算ROIを計算する。"""
    history_by_horse: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for horse in horse_history:
        if horse.get("horse_id") and horse.get("finish") is not None:
            history_by_horse[str(horse["horse_id"])].append(horse)
    for members in history_by_horse.values():
        members.sort(key=lambda row: (str(row["date"]), str(row.get("source_race_id") or "")))

    rows: list[dict[str, Any]] = []
    for signal in signals:
        horse_id = str(signal.get("horse_id") or "")
        signal_date = str(signal["signal_date"])
        valid_until = str(signal.get("valid_until") or "")
        valid_until_compact = valid_until.replace("-", "")
        candidates = [
            row
            for row in history_by_horse.get(horse_id, [])
            if str(row["date"]) > signal_date
            and (not valid_until_compact or str(row["date"]) <= valid_until_compact)
        ]
        next_start = candidates[0] if candidates else None
        if next_start is not None:
            odds = next_start.get("odds")
            finish = next_start.get("finish")
            has_odds = odds is not None and not pd.isna(odds)
            stake = 100 if has_odds else None
            win_return = round(float(odds) * 100, 1) if has_odds and int(finish) == 1 else (0.0 if has_odds else None)
            roi = round(win_return / stake * 100, 1) if stake else None
            status = "SETTLED" if has_odds else "SETTLED_NO_ODDS"
        else:
            odds = finish = stake = win_return = roi = None
            status = "EXPIRED" if valid_until_compact and as_of_date > valid_until_compact else "PENDING"
        rows.append(
            {
                "signal_type": signal["signal_type"],
                "signal_date": signal_date,
                "signal_source_race_id": signal["signal_source_race_id"],
                "horse_id": horse_id,
                "horse_name": signal["horse_name"],
                "valid_until": valid_until or None,
                "status": status,
                "next_date": next_start["date"] if next_start else None,
                "next_source_race_id": next_start.get("source_race_id") if next_start else None,
                "next_racecourse": next_start.get("racecourse") if next_start else None,
                "next_race_no": next_start.get("race_no") if next_start else None,
                "next_finish": int(finish) if finish is not None else None,
                "next_popularity": next_start.get("popularity") if next_start else None,
                "next_odds": float(odds) if odds is not None and not pd.isna(odds) else None,
                "win_stake": stake,
                "win_return": win_return,
                "win_roi_pct": roi,
            }
        )
    return rows


def summarize_next_start_outcomes(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """注目・危険シグナル別に、次走成績と単勝仮想回収率を集計する。"""
    summaries: list[dict[str, Any]] = []
    for signal_type in ("WATCH", "DANGER"):
        members = [row for row in rows if row["signal_type"] == signal_type]
        settled = [row for row in members if row["status"] == "SETTLED"]
        stake = sum(float(row["win_stake"]) for row in settled if row["win_stake"] is not None)
        returns = sum(float(row["win_return"]) for row in settled if row["win_return"] is not None)
        summaries.append(
            {
                "signal_type": signal_type,
                "signals": len(members),
                "settled": len(settled),
                "wins": sum(1 for row in settled if row["next_finish"] == 1),
                "average_popularity": round(
                    statistics.mean(float(row["next_popularity"]) for row in settled if row["next_popularity"] is not None),
                    2,
                ) if any(row["next_popularity"] is not None for row in settled) else None,
                "win_stake": stake,
                "win_return": returns,
                "win_roi_pct": round(returns / stake * 100, 1) if stake else None,
            }
        )
    return summaries


def model_review_rows(race_date: str) -> list[dict[str, Any]]:
    """即時採用せず、再現性検証が必要な改善仮説を保存する。"""
    return [
        {
            "issue_key": f"{race_date}:PIPELINE:PRE_RACE_ARTIFACT",
            "date": race_date,
            "issue_type": "データ取得・運用障害",
            "racecourse": "全場",
            "race_no": None,
            "problem_description": "真正な事前予想・T-5成果物が保存されておらず、結果との比較が不能",
            "current_model_behavior": "予想順位・score・購入判断をNULLで保存",
            "improvement_hypothesis": "Morning完了条件に予想成果物の存在確認を追加すると欠落開催を検知できる",
            "required_data": "Morningログ、予想成果物、T-5成果物、tracking DB",
            "backtest_method": "一時環境でMorningを再現し、成果物生成と失敗検知を確認",
            "priority": "HIGH",
            "sample_size_required": "1回の再現テスト＋次回開催のSHADOW確認",
            "overfitting_risk": "LOW",
            "status": "NEW",
        },
        {
            "issue_key": f"{race_date}:FEATURE:TRACK_PATH",
            "date": race_date,
            "issue_type": "新特徴量候補",
            "racecourse": "全場",
            "race_no": None,
            "problem_description": "枠番と通過順位だけでは実際の内外進路・距離ロスを判定できない",
            "current_model_behavior": "内外バイアスと不利を欠損扱い",
            "improvement_hypothesis": "4角と直線の走行レーンを映像由来で構造化すると再現可能なバイアス判定が可能",
            "required_data": "走行レーン、コーナー距離ロス、直線進路",
            "backtest_method": "複数開催で人手二重ラベルし、次走成績との関連を検証",
            "priority": "MEDIUM",
            "sample_size_required": "10開催以上",
            "overfitting_risk": "MEDIUM",
            "status": "NEW",
        },
        {
            "issue_key": f"{race_date}:BASELINE:CLOCK",
            "date": race_date,
            "issue_type": "基準時計改善候補",
            "racecourse": "全場",
            "race_no": None,
            "problem_description": "同距離・同馬場・同クラスの標本が少ない条件がある",
            "current_model_behavior": "標本不足時は距離×馬場、距離×クラス、距離の順でフォールバック",
            "improvement_hypothesis": "季節・コース区分・斤量・年齢条件を含む階層モデルで標準時計の不確実性を推定できる",
            "required_data": "長期勝ち時計、馬場、クラス、コース区分、気象",
            "backtest_method": "時系列分割で標準時計誤差と校正を比較",
            "priority": "MEDIUM",
            "sample_size_required": "各主要条件30レース以上",
            "overfitting_risk": "MEDIUM",
            "status": "NEW",
        },
    ]


def ensure_columns(connection: sqlite3.Connection, table: str, columns: dict[str, str]) -> None:
    """既存SQLite表へ不足列だけを追加し、保存済みデータを維持する。"""
    existing = {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}
    for name, definition in columns.items():
        if name not in existing:
            connection.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")


def create_database(path: Path) -> sqlite3.Connection:
    """回顧専用SQLiteを作成し、重複防止キーを持つ表を準備する。"""
    connection = sqlite3.connect(path)
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS track_daily_bias (
            date TEXT NOT NULL, racecourse TEXT NOT NULL, surface TEXT NOT NULL,
            distance_group TEXT NOT NULL, going TEXT, race_count INTEGER,
            clock_bias TEXT, clock_bias_score REAL, clock_diff_sec REAL,
            clock_diff_per_1000m REAL,
            inside_outside_bias TEXT, inside_outside_score REAL,
            pace_style_bias TEXT, pace_style_score REAL,
            front_runner_score REAL, stalker_score REAL, midpack_score REAL,
            closer_score REAL, deep_closer_score REAL,
            inside_win_rate REAL, outside_win_rate REAL,
            front_win_rate REAL, front_place_rate REAL,
            closer_win_rate REAL, closer_place_rate REAL,
            early_race_bias TEXT, middle_race_bias TEXT, late_race_bias TEXT,
            weather TEXT, rainfall REAL, wind_direction TEXT, wind_speed REAL,
            confidence TEXT, notes TEXT, style_stats_json TEXT, created_at TEXT,
            PRIMARY KEY (date, racecourse, surface, distance_group)
        );
        CREATE TABLE IF NOT EXISTS race_review (
            date TEXT NOT NULL, racecourse TEXT NOT NULL, race_no INTEGER NOT NULL,
            source_race_id TEXT, race_name TEXT, surface TEXT, distance INTEGER,
            distance_group TEXT, going TEXT, class_name TEXT, field_size INTEGER,
            weather TEXT, start_time TEXT, winner TEXT, second TEXT, third TEXT,
            winning_time TEXT, pace_type TEXT, early3f REAL, late3f_race REAL,
            first1000m TEXT, lap_times TEXT, race_level TEXT,
            clock_rating TEXT, clock_diff_sec REAL, clock_diff_per_1000m REAL, clock_reference_sec REAL,
            clock_reference_n INTEGER, clock_reference_basis TEXT, pace_rating TEXT,
            inside_outside_bias TEXT, running_style_bias TEXT,
            prediction_top1 TEXT, prediction_top1_finish INTEGER,
            prediction_top3_hit INTEGER, prediction_top5_hit INTEGER,
            winner_prediction_rank INTEGER, payout TEXT, review_summary TEXT,
            prediction_error_type TEXT, bet_or_skip_review TEXT, source_url TEXT,
            created_at TEXT, PRIMARY KEY (date, racecourse, race_no)
        );
        CREATE TABLE IF NOT EXISTS horse_review (
            date TEXT NOT NULL, racecourse TEXT NOT NULL, race_no INTEGER NOT NULL,
            source_race_id TEXT, horse_id TEXT, horse_name TEXT NOT NULL, horse_no INTEGER, frame_no INTEGER,
            finish INTEGER, finish_status TEXT, popularity INTEGER, odds REAL,
            finish_time TEXT, margin TEXT, last3f REAL, passing_order TEXT,
            last_corner_position INTEGER, carried_weight REAL, jockey TEXT,
            trainer TEXT, body_weight TEXT, prediction_rank INTEGER,
            prediction_score REAL, running_style TEXT, actual_performance TEXT,
            trip_trouble TEXT, pace_advantage INTEGER, pace_disadvantage INTEGER,
            track_bias_advantage INTEGER, track_bias_disadvantage INTEGER,
            inside_outside_advantage INTEGER, distance_loss REAL,
            late_speed_rating REAL, opponent_strength TEXT, race_level TEXT,
            next_start_watch_score REAL, next_start_watch_reason TEXT,
            next_start_target_condition TEXT, next_start_risk TEXT, created_at TEXT,
            PRIMARY KEY (date, racecourse, race_no, horse_name)
        );
        CREATE TABLE IF NOT EXISTS next_start_watch (
            date TEXT NOT NULL, source_race_id TEXT, horse_id TEXT, horse_name TEXT NOT NULL, racecourse TEXT NOT NULL,
            race_no INTEGER NOT NULL, finish INTEGER, popularity INTEGER,
            watch_rank INTEGER, watch_reason TEXT, trouble_type TEXT,
            hidden_strength TEXT, preferred_surface TEXT, preferred_distance TEXT,
            preferred_course TEXT, preferred_going TEXT,
            expected_market_mispricing TEXT, avoid_if_favorite INTEGER,
            valid_until TEXT, created_at TEXT,
            PRIMARY KEY (date, racecourse, race_no, horse_name)
        );
        CREATE TABLE IF NOT EXISTS dangerous_next_start (
            date TEXT NOT NULL, source_race_id TEXT, horse_id TEXT, horse_name TEXT NOT NULL, racecourse TEXT NOT NULL,
            race_no INTEGER NOT NULL, finish INTEGER, popularity INTEGER,
            advantage_type TEXT, pace_advantage INTEGER,
            track_bias_advantage INTEGER, trip_advantage INTEGER,
            danger_reason TEXT, danger_condition TEXT, valid_until TEXT, created_at TEXT,
            PRIMARY KEY (date, racecourse, race_no, horse_name)
        );
        CREATE TABLE IF NOT EXISTS model_review (
            issue_key TEXT PRIMARY KEY, date TEXT, issue_type TEXT,
            racecourse TEXT, race_no INTEGER, problem_description TEXT,
            current_model_behavior TEXT, improvement_hypothesis TEXT,
            required_data TEXT, backtest_method TEXT, priority TEXT,
            sample_size_required TEXT, overfitting_risk TEXT, status TEXT,
            created_at TEXT
        );
        CREATE TABLE IF NOT EXISTS horse_master (
            horse_id TEXT PRIMARY KEY, horse_name TEXT NOT NULL,
            first_seen_date TEXT NOT NULL, last_seen_date TEXT NOT NULL,
            created_at TEXT, updated_at TEXT
        );
        CREATE TABLE IF NOT EXISTS track_condition_stats (
            date TEXT NOT NULL, racecourse TEXT NOT NULL, surface TEXT NOT NULL,
            distance_group TEXT NOT NULL, going TEXT NOT NULL, class_name TEXT NOT NULL,
            race_count INTEGER, clock_diff_per_1000m_mean REAL,
            clock_diff_per_1000m_median REAL, clock_diff_per_1000m_std REAL,
            clock_diff_per_1000m_min REAL, clock_diff_per_1000m_max REAL,
            clock_bias TEXT, confidence TEXT, created_at TEXT,
            PRIMARY KEY (date, racecourse, surface, distance_group, going, class_name)
        );
        CREATE TABLE IF NOT EXISTS next_start_outcome (
            signal_type TEXT NOT NULL, signal_date TEXT NOT NULL,
            signal_source_race_id TEXT NOT NULL, horse_id TEXT NOT NULL,
            horse_name TEXT, valid_until TEXT, status TEXT,
            next_date TEXT, next_source_race_id TEXT, next_racecourse TEXT,
            next_race_no INTEGER, next_finish INTEGER, next_popularity INTEGER,
            next_odds REAL, win_stake REAL, win_return REAL, win_roi_pct REAL,
            created_at TEXT, updated_at TEXT,
            PRIMARY KEY (signal_type, signal_date, signal_source_race_id, horse_id)
        );
        """
    )
    ensure_columns(connection, "track_daily_bias", {"clock_diff_per_1000m": "REAL"})
    ensure_columns(connection, "race_review", {"clock_diff_per_1000m": "REAL"})
    ensure_columns(connection, "horse_review", {"source_race_id": "TEXT", "horse_id": "TEXT"})
    ensure_columns(connection, "next_start_watch", {"source_race_id": "TEXT", "horse_id": "TEXT"})
    ensure_columns(
        connection,
        "dangerous_next_start",
        {"source_race_id": "TEXT", "horse_id": "TEXT", "valid_until": "TEXT"},
    )
    connection.executescript(
        """
        CREATE INDEX IF NOT EXISTS idx_horse_review_horse_date
            ON horse_review (horse_id, date);
        CREATE INDEX IF NOT EXISTS idx_horse_review_source_race
            ON horse_review (source_race_id);
        CREATE INDEX IF NOT EXISTS idx_next_start_outcome_next_date
            ON next_start_outcome (next_date, signal_type);
        """
    )
    return connection


def upsert_rows(connection: sqlite3.Connection, table: str, rows: list[dict[str, Any]], keys: list[str], created_at: str) -> int:
    """同一キーを二重登録せず、再実行時はより新しい生成結果で更新する。"""
    if not rows:
        return 0
    columns = list(rows[0].keys()) + ["created_at"]
    placeholders = ",".join("?" for _ in columns)
    update_columns = [c for c in columns if c not in keys]
    update_sql = ",".join(f"{c}=excluded.{c}" for c in update_columns)
    sql = (
        f"INSERT INTO {table} ({','.join(columns)}) VALUES ({placeholders}) "
        f"ON CONFLICT ({','.join(keys)}) DO UPDATE SET {update_sql}"
    )
    values = [[row.get(column) if column != "created_at" else created_at for column in columns] for row in rows]
    connection.executemany(sql, values)
    return len(rows)


def upsert_horse_master(connection: sqlite3.Connection, rows: list[dict[str, Any]], timestamp: str) -> int:
    """馬マスタをhorse_idで更新し、初回・最終観測日を維持する。"""
    sql = """
        INSERT INTO horse_master
            (horse_id, horse_name, first_seen_date, last_seen_date, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT (horse_id) DO UPDATE SET
            horse_name=excluded.horse_name,
            first_seen_date=MIN(horse_master.first_seen_date, excluded.first_seen_date),
            last_seen_date=MAX(horse_master.last_seen_date, excluded.last_seen_date),
            updated_at=excluded.updated_at
    """
    connection.executemany(
        sql,
        [
            (
                row["horse_id"],
                row["horse_name"],
                row["first_seen_date"],
                row["last_seen_date"],
                timestamp,
                timestamp,
            )
            for row in rows
        ],
    )
    return len(rows)


def refresh_next_start_outcomes(
    connection: sqlite3.Connection,
    timestamp: str,
    as_of_date: str,
) -> list[dict[str, Any]]:
    """保存済み全シグナルを最新のhorse_reviewへ結び直し、未解決分も継続追跡する。"""
    connection.row_factory = sqlite3.Row
    signals = [
        dict(row)
        for row in connection.execute(
            """
            SELECT 'WATCH' AS signal_type, date AS signal_date,
                   source_race_id AS signal_source_race_id, horse_id, horse_name, valid_until
            FROM next_start_watch
            WHERE horse_id IS NOT NULL AND source_race_id IS NOT NULL
            UNION ALL
            SELECT 'DANGER' AS signal_type, date AS signal_date,
                   source_race_id AS signal_source_race_id, horse_id, horse_name, valid_until
            FROM dangerous_next_start
            WHERE horse_id IS NOT NULL AND source_race_id IS NOT NULL
            """
        )
    ]
    history = [
        dict(row)
        for row in connection.execute(
            """
            SELECT date, source_race_id, horse_id, horse_name, racecourse, race_no,
                   finish, popularity, odds
            FROM horse_review
            WHERE horse_id IS NOT NULL
            """
        )
    ]
    outcomes = calculate_next_start_outcomes(signals, history, as_of_date=as_of_date)
    sql = """
        INSERT INTO next_start_outcome
            (signal_type, signal_date, signal_source_race_id, horse_id, horse_name,
             valid_until, status, next_date, next_source_race_id, next_racecourse,
             next_race_no, next_finish, next_popularity, next_odds, win_stake,
             win_return, win_roi_pct, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT (signal_type, signal_date, signal_source_race_id, horse_id)
        DO UPDATE SET
            horse_name=excluded.horse_name,
            valid_until=excluded.valid_until,
            status=excluded.status,
            next_date=excluded.next_date,
            next_source_race_id=excluded.next_source_race_id,
            next_racecourse=excluded.next_racecourse,
            next_race_no=excluded.next_race_no,
            next_finish=excluded.next_finish,
            next_popularity=excluded.next_popularity,
            next_odds=excluded.next_odds,
            win_stake=excluded.win_stake,
            win_return=excluded.win_return,
            win_roi_pct=excluded.win_roi_pct,
            updated_at=excluded.updated_at
    """
    connection.executemany(
        sql,
        [
            (
                row["signal_type"], row["signal_date"], row["signal_source_race_id"],
                row["horse_id"], row["horse_name"], row["valid_until"], row["status"],
                row["next_date"], row["next_source_race_id"], row["next_racecourse"],
                row["next_race_no"], row["next_finish"], row["next_popularity"],
                row["next_odds"], row["win_stake"], row["win_return"],
                row["win_roi_pct"], timestamp, timestamp,
            )
            for row in outcomes
        ],
    )
    return outcomes


def markdown_report(
    races: list[dict[str, Any]],
    horses: list[dict[str, Any]],
    biases: list[dict[str, Any]],
    condition_stats: list[dict[str, Any]],
    watches: list[dict[str, Any]],
    dangers: list[dict[str, Any]],
    outcome_summaries: list[dict[str, Any]],
    models: list[dict[str, Any]],
    race_date: str,
) -> str:
    """指定構成に沿ったMarkdown回顧レポートを生成する。"""
    longshots = [h for h in horses if h["finish"] is not None and h["finish"] <= 3 and (h["popularity"] or 0) >= 6]
    favorite_failures = [h for h in horses if (h["popularity"] or 99) <= 3 and (h["finish"] or 99) >= 4]
    high_uncertainty = [r for r in races if r["class_name"] == "新馬" or r["surface"] == "障害"]
    date_label = datetime.strptime(race_date, "%Y%m%d").strftime("%Y-%m-%d")
    course_counts = {
        course: sum(1 for race in races if race["racecourse"] == course)
        for course in sorted({race["racecourse"] for race in races})
    }
    course_summary = "、".join(f"{course}{count}R" for course, count in course_counts.items())
    special_note = "中山は開催中止のため対象外。" if race_date == "20260921" else ""
    lines = [
        f"# {date_label} JRA全レース回顧",
        "",
        f"> {course_summary}を対象とした。{special_note}真正な事前予想成果物がないため予想比較欄はNULLとし、人気を予想順位へ代用していない。",
        "",
        "## 1. 開催全体総括",
        "",
        f"{course_summary}、計{len(races)}R（障害{sum(r['surface']=='障害' for r in races)}、芝{sum(r['surface']=='芝' for r in races)}、ダート{sum(r['surface']=='ダート' for r in races)}）を収録した。全着順・上がり・通過順位・ラップ・払戻を構造化し、映像が必要な不利と内外進路は推測していない。",
        "",
        "## 2. 競馬場別・芝馬場分析",
        "",
    ]
    for bias in [b for b in biases if b["surface"] == "芝"]:
        lines.extend(
            [
                f"### {bias['racecourse']}芝",
                "",
                f"時計：{bias['clock_bias']}（1000m当たり標準比平均 {bias['clock_diff_per_1000m']}秒、生時計差平均 {bias['clock_diff_sec']}秒）",
                "",
                "内外：データなし（実走進路を未取得）",
                "",
                f"脚質：{bias['pace_style_bias']}（confidence: {bias['confidence']}）",
                "",
                f"時間帯：{bias['early_race_bias']} / {bias['middle_race_bias']} / {bias['late_race_bias']}",
                "",
                f"総評：対象{bias['race_count']}R。時計は既存結果ブックの同距離条件中央値と比較。脚質は最終コーナー位置比による機械集計であり、因果的な馬場有利とは断定しない。",
                "",
            ]
        )
    lines.extend(["## 3. 競馬場別・ダート分析", ""])
    for bias in [b for b in biases if b["surface"] == "ダート"]:
        lines.extend(
            [
                f"### {bias['racecourse']}ダート",
                "",
                f"時計：{bias['clock_bias']}（1000m当たり標準比平均 {bias['clock_diff_per_1000m']}秒、生時計差平均 {bias['clock_diff_sec']}秒）",
                "",
                "内外：データなし（実走進路を未取得）",
                "",
                f"脚質：{bias['pace_style_bias']}（confidence: {bias['confidence']}）",
                "",
                f"時間帯：{bias['early_race_bias']} / {bias['middle_race_bias']} / {bias['late_race_bias']}",
                "",
                f"総評：対象{bias['race_count']}R。脚質集計は出走構成の影響を含むため、単日だけでモデルへ採用しない。",
                "",
            ]
        )
    lines.extend(["## 4. 時間帯による変化", "", "前半・中盤・後半の勝ち馬脚質集計は上記に記録した。芝・ダートとも少数標本のため、降雨・乾燥・風を変化原因として断定しない。", ""])
    lines.extend(["## 5. 距離別特徴", ""])
    for racecourse in sorted(course_counts):
        for surface in ("芝", "ダート"):
            for group in ("短距離", "マイル", "中距離", "長距離"):
                subset = [r for r in races if r["racecourse"] == racecourse and r["surface"] == surface and r["distance_group"] == group]
                if subset:
                    pace_counts: dict[str, int] = {}
                    for race in subset:
                        pace_counts[race["pace_type"]] = pace_counts.get(race["pace_type"], 0) + 1
                    lines.append(f"- {racecourse}{surface}{group}：{len(subset)}R、ペース {json.dumps(pace_counts, ensure_ascii=False)}。少数区分は傾向断定なし。")
    lines.extend(["", "### 距離・馬場・クラス別の正規化時計", ""])
    for row in condition_stats:
        lines.append(
            f"- {row['racecourse']}{row['surface']}{row['distance_group']}・{row['going']}・{row['class_name']}："
            f"1000m当たり標準差 平均{row['clock_diff_per_1000m_mean']:+.4f}秒、"
            f"中央値{row['clock_diff_per_1000m_median']:+.4f}秒、n={row['race_count']}、"
            f"confidence={row['confidence']}。"
        )
    lines.extend(["", "## 6. 重要レース回顧", ""])
    important_keys = {(h["racecourse"], h["race_no"]) for h in longshots + favorite_failures}
    important_keys.update((h["racecourse"], h["race_no"]) for h in sorted(horses, key=lambda x: -x["next_start_watch_score"])[:5])
    for race in [r for r in races if (r["racecourse"], r["race_no"]) in important_keys][:10]:
        lines.append(f"### {race['racecourse']}{race['race_no']}R {race['race_name']}")
        lines.append("")
        lines.append(race["review_summary"])
        lines.append("")
    lines.extend(["## 7. 次走注目馬 TOP10", "", "|順位|競馬場R|馬名|今回着順|人気|理由|次走狙い条件|", "|---:|---|---|---:|---:|---|---|"])
    for row in watches:
        lines.append(f"|{row['watch_rank']}|{row['racecourse']}{row['race_no']}R|{row['horse_name']}|{row['finish']}|{row['popularity']}|{row['watch_reason']}|今回より展開適合度が高い条件|")
    if not watches:
        lines.append("|–|–|該当なし|–|–|定量基準を満たす馬なし|–|")
    lines.extend(["", "## 8. 次走危険馬", "", "|競馬場R|馬名|今回着順|人気|危険理由|", "|---|---|---:|---:|---|"])
    for row in dangers:
        lines.append(f"|{row['racecourse']}{row['race_no']}R|{row['horse_name']}|{row['finish']}|{row['popularity']}|{row['danger_reason']}。{row['danger_condition']}|")
    if not dangers:
        lines.append("|–|該当なし|–|–|客観的な展開利基準を満たす馬なし|")
    lines.extend(["", "## 9. 人気薄激走", ""])
    for horse in longshots:
        signal = f"上がり{horse['late_speed_rating']}位" if horse["late_speed_rating"] else "上がり順位データなし"
        reproducibility = "再検証候補" if horse["late_speed_rating"] and horse["late_speed_rating"] <= 3 else "偶然性を排除できない"
        lines.append(f"- {horse['racecourse']}{horse['race_no']}R {horse['horse_name']}（{horse['popularity']}人気{horse['finish']}着）：{signal}。分類={reproducibility}。人気薄の理由はデータなし。")
    if not longshots:
        lines.append("- 該当なし。")
    lines.extend(["", "## 10. 人気馬凡走", ""])
    for horse in favorite_failures:
        cause = f"{horse['running_style']}に{('展開不利' if horse['pace_disadvantage'] else '明確な定量原因なし')}"
        lines.append(f"- {horse['racecourse']}{horse['race_no']}R {horse['horse_name']}（{horse['popularity']}人気{horse['finish']}着）：{cause}。映像由来の不利は未判定。")
    if not favorite_failures:
        lines.append("- 該当なし。")
    lines.extend(
        [
            "",
            "## 11. 予想モデルとのズレ",
            "",
            "当日予想順位・score・T-5判断は保存されていない。結果取得後の再予想は比較に使わず、人気順位による代用もしていない。",
            "",
            "## 12. 買わない方がよかったレース",
            "",
        ]
    )
    for race in high_uncertainty:
        reason = "初出走馬中心で事前情報が少ない" if race["class_name"] == "新馬" else "障害で通常平地モデルの適用外"
        lines.append(f"- {race['racecourse']}{race['race_no']}R {race['race_name']}：{reason}。")
    lines.extend(["", "全レースについて当日予想が存在しないため、購入可否の事後判定は行わない。", "", "## 13. 新特徴量・改善仮説", ""])
    for model in models:
        lines.append(f"- [{model['priority']}] {model['improvement_hypothesis']}（status={model['status']}、即時採用なし）")
    lines.extend(
        [
            "",
            "## 14. DBへ新規保存した重要事項",
            "",
            f"race_review={len(races)}件、horse_review={len(horses)}件、track_daily_bias={len(biases)}件、track_condition_stats={len(condition_stats)}件、next_start_watch={len(watches)}件、dangerous_next_start={len(dangers)}件、model_review={len(models)}件。",
            "",
            "### 次走シグナル追跡（単勝100円仮想購入）",
            "",
        ]
    )
    for summary in outcome_summaries:
        roi_text = f"{summary['win_roi_pct']:.1f}%" if summary["win_roi_pct"] is not None else "未確定"
        lines.append(
            f"- {summary['signal_type']}：候補{summary['signals']}件、次走確定{summary['settled']}件、"
            f"1着{summary['wins']}件、単勝回収率={roi_text}。"
        )
    lines.extend(
        [
            "",
            "## 15. 次回開催で使うポイント",
            "",
            "1. 予想成果物の存在をMorning成功条件として機械確認する。",
            "2. 予想がない開催は、人気をモデル順位へ代用しない。",
            "3. 次走注目は上がり順位とペース不利が同時に確認できる馬を優先する。",
            "4. 好走馬はペース利が再現しない条件で過信しない。",
            "5. 新馬・障害は情報不足として見送り基準を強める。",
            "6. 内外バイアスは枠番だけで判定しない。",
            "7. 時計評価は同距離・同馬場・同クラスの標本数を併記する。",
            "8. 単日脚質集計は因果的な馬場傾向とみなさない。",
            "9. 風向・風速が欠損した日は風の影響を断定しない。",
            "10. 改善仮説はバックテスト完了まで本番モデルへ入れない。",
            "",
            f"## 全{len(races)}R一覧",
            "",
            "|競馬場|R|レース|条件|馬場|頭数|1着|2着|3着|ペース|時計評価|",
            "|---|---:|---|---|---|---:|---|---|---|---|---|",
        ]
    )
    for race in races:
        condition = f"{race['surface']}{race['distance']}m"
        lines.append(f"|{race['racecourse']}|{race['race_no']}|{race['race_name']}|{condition}|{race['going']}|{race['field_size']}|{race['winner']}|{race['second']}|{race['third']}|{race['pace_type']}|{race['clock_rating']}|")
    lines.extend(
        [
            "",
            "## データ品質・判定ルール",
            "",
            "- actual_performance=A：4～12着で上がり3位以内、かつペース不利等を加点した注目スコア50以上。",
            "- C/D：3着以内でペース利あり。極端ペース時のみD、それ以外はC。その他はB。",
            "- 内外・接触・前詰まり・大外回し・風は、実測または映像データがないためNULL／データなし。",
            "- 時計標準は既存 `racedata_results.xlsx` の同競馬場勝ち時計中央値。同距離×馬場×クラスを優先し、n<3なら段階的に条件を緩和。",
            "- 予想比較列は、真正な事前予想がないためNULL。",
            "",
            "## 出典",
            "",
            f"- [JRA開催日程]({jra_schedule_url(race_date)})",
            "- 各レース結果：netkeiba（各race_review行のsource_urlに保存）",
        ]
    )
    if race_date == "20260921":
        lines.append("- [JRA中山中止・代替競馬告知](https://www.jra.go.jp/news/202609/092105.html)")
    return "\n".join(lines) + "\n"


def validate_scraped_payload(
    races: list[dict[str, Any]],
    horses: list[dict[str, Any]],
    official_counts: dict[str, int],
) -> None:
    """DB・CSV作成前に、公式件数・レース一意性・馬ID完全性を検証する。"""
    race_ids = [str(race["source_race_id"]) for race in races]
    verify_official_race_counts(race_ids, official_counts)
    if len(race_ids) != len(set(race_ids)):
        raise RuntimeError("取得結果に重複レースIDがあるため保存を中止します")
    missing_horse_ids = [
        f"{horse.get('source_race_id')}:{horse.get('horse_no')}:{horse.get('horse_name')}"
        for horse in horses
        if not horse.get("horse_id")
    ]
    if missing_horse_ids:
        preview = ", ".join(missing_horse_ids[:5])
        raise RuntimeError(
            f"horse_id欠損が{len(missing_horse_ids)}件あるため保存を中止します: {preview}"
        )


def main() -> None:
    """対象日の全レースを取得し、CSV・SQLite・Markdownへ一括保存する。"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--race-date", default=DEFAULT_RACE_DATE)
    args = parser.parse_args()
    project_root = args.project_root.resolve()
    review_root = project_root / "races_review"
    dated_dir = review_root / args.race_date[:4] / args.race_date[:6]

    history_path = project_root / "data" / "master" / "racedata_results.xlsx"
    if not history_path.exists():
        raise FileNotFoundError("過去結果ブックがありません")

    official_counts, race_ids = fetch_official_race_program(args.race_date)
    verify_official_race_counts(race_ids, official_counts)

    reference = load_historical_clock_reference(history_path)
    entry_horse_ids = load_entry_horse_id_map(project_root, args.race_date)
    races: list[dict[str, Any]] = []
    horses: list[dict[str, Any]] = []
    for race_id in race_ids:
        race, runners = scrape_race(race_id, args.race_date, reference, entry_horse_ids)
        races.append(race)
        horses.extend(runners)
        print(f"取得: {race['racecourse']}{race['race_no']}R {race['race_name']} {len(runners)}頭")

    validate_scraped_payload(races, horses, official_counts)

    enrich_race_summaries(races, horses)
    biases = build_track_bias(races, horses)
    condition_stats = build_track_condition_stats(races)
    watches = build_watch_rows(horses, args.race_date)
    dangers = build_danger_rows(horses, races, args.race_date)
    horse_master = build_horse_master_rows(horses)
    models = model_review_rows(args.race_date)

    now = datetime.now(ZoneInfo("Asia/Tokyo")).isoformat(timespec="seconds")
    database_path = review_root / "race_review.db"
    connection = create_database(database_path)
    counts = {
        "track_daily_bias": upsert_rows(connection, "track_daily_bias", biases, ["date", "racecourse", "surface", "distance_group"], now),
        "track_condition_stats": upsert_rows(connection, "track_condition_stats", condition_stats, ["date", "racecourse", "surface", "distance_group", "going", "class_name"], now),
        "race_review": upsert_rows(connection, "race_review", races, ["date", "racecourse", "race_no"], now),
        "horse_review": upsert_rows(connection, "horse_review", horses, ["date", "racecourse", "race_no", "horse_name"], now),
        "next_start_watch": upsert_rows(connection, "next_start_watch", watches, ["date", "racecourse", "race_no", "horse_name"], now),
        "dangerous_next_start": upsert_rows(connection, "dangerous_next_start", dangers, ["date", "racecourse", "race_no", "horse_name"], now),
        "horse_master": upsert_horse_master(connection, horse_master, now),
        "model_review": upsert_rows(connection, "model_review", models, ["issue_key"], now),
    }
    outcomes = refresh_next_start_outcomes(connection, now, args.race_date)
    outcome_summaries = summarize_next_start_outcomes(outcomes)
    counts["next_start_outcome"] = len(outcomes)
    connection.commit()
    connection.close()

    dated_dir.mkdir(parents=True, exist_ok=True)
    reference.to_csv(dated_dir / f"historical_clock_reference_{args.race_date}.csv", index=False, encoding="utf-8-sig")
    csv_outputs = {
        f"track_daily_bias_{args.race_date}.csv": biases,
        f"track_condition_stats_{args.race_date}.csv": condition_stats,
        f"race_review_{args.race_date}.csv": races,
        f"horse_review_{args.race_date}.csv": horses,
        f"next_start_watch_{args.race_date}.csv": watches,
        f"dangerous_next_start_{args.race_date}.csv": dangers,
        f"next_start_outcome_{args.race_date}.csv": outcomes,
        f"model_review_{args.race_date}.csv": models,
    }
    for filename, rows in csv_outputs.items():
        pd.DataFrame(rows).to_csv(dated_dir / filename, index=False, encoding="utf-8-sig")

    report_path = dated_dir / f"race_review_{args.race_date}.md"
    report_path.write_text(
        markdown_report(
            races,
            horses,
            biases,
            condition_stats,
            watches,
            dangers,
            outcome_summaries,
            models,
            args.race_date,
        ),
        encoding="utf-8",
    )

    print("\n=== 回顧処理サマリー ===")
    print(f"Markdown: {report_path}")
    print("CSV: " + ", ".join(str(dated_dir / name) for name in csv_outputs))
    print(f"DB: {database_path}")
    print("DB保存件数: " + json.dumps(counts, ensure_ascii=False))
    print(f"新規次走注目馬数: {len(watches)}")
    print("次走追跡: " + json.dumps(outcome_summaries, ensure_ascii=False))
    print(f"新規モデル改善仮説数: {len(models)}")
    for bias in biases:
        print(f"{bias['racecourse']}{bias['surface']}: 時計={bias['clock_bias']} 脚質={bias['pace_style_bias']} 内外=データなし confidence={bias['confidence']}")


if __name__ == "__main__":
    main()
